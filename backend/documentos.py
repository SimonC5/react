"""Generación de los documentos descargables: reportes PDF/Excel y facturas PDF.

Las librerías se importan dentro de cada función para que la API siga
levantando aunque falten; en ese caso el endpoint responde 503 con la
instrucción de instalar ``backend/requirements.txt``.
"""

from io import BytesIO
from typing import Any, Optional

from fastapi import HTTPException

EMPRESA = 'SimonC Realidad Virtual'


def _dependencia_faltante(nombre: str, paquete: str) -> HTTPException:
    return HTTPException(
        status_code=503,
        detail=f'{nombre} no está disponible. Instala las dependencias con "pip install -r backend/requirements.txt" ({paquete}).',
    )


def formato_moneda(valor: Any) -> str:
    try:
        numero = float(valor or 0)
    except (TypeError, ValueError):
        numero = 0.0
    return f'$ {numero:,.2f}'


def construir_pdf(
    titulo: str,
    subtitulo: str,
    encabezados: list[str],
    filas: list[list[str]],
    totales: Optional[list[tuple[str, str]]] = None,
    notas: Optional[list[str]] = None,
) -> bytes:
    """Arma un PDF tabular con encabezado, tabla y bloque de totales."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise _dependencia_faltante('La generación de PDF', 'reportlab') from exc

    buffer = BytesIO()
    ancha = len(encabezados) > 5
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4) if ancha else A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=titulo,
        author=EMPRESA,
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph(f'<b>{EMPRESA}</b>', estilos['Title']),
        Paragraph(titulo, estilos['Heading2']),
        Paragraph(subtitulo, estilos['Normal']),
        Spacer(1, 8 * mm),
    ]

    datos = [encabezados] + (filas or [['Sin registros para los filtros seleccionados.'] + [''] * (len(encabezados) - 1)])
    tabla = Table(datos, repeatRows=1, hAlign='LEFT')
    tabla.setStyle(
        TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#cbd5f5')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f1f5f9')]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ])
    )
    elementos.append(tabla)

    if totales:
        elementos.append(Spacer(1, 6 * mm))
        tabla_totales = Table([[etiqueta, valor] for etiqueta, valor in totales], hAlign='RIGHT')
        tabla_totales.setStyle(
            TableStyle([
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
                ('LINEABOVE', (0, 0), (-1, 0), 0.6, colors.HexColor('#0f172a')),
            ])
        )
        elementos.append(tabla_totales)

    for nota in notas or []:
        elementos.append(Spacer(1, 4 * mm))
        elementos.append(Paragraph(nota, estilos['Italic']))

    documento.build(elementos)
    return buffer.getvalue()


def construir_xlsx(
    titulo: str,
    subtitulo: str,
    encabezados: list[str],
    filas: list[list[Any]],
    totales: Optional[list[tuple[str, Any]]] = None,
) -> bytes:
    """Arma un libro de Excel con la misma información del reporte en PDF."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise _dependencia_faltante('La exportación a Excel', 'openpyxl') from exc

    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Reporte'

    hoja.append([EMPRESA])
    hoja['A1'].font = Font(size=14, bold=True)
    hoja.append([titulo])
    hoja['A2'].font = Font(size=12, bold=True)
    hoja.append([subtitulo])
    hoja.append([])

    fila_encabezado = hoja.max_row + 1
    hoja.append(encabezados)
    relleno = PatternFill('solid', start_color='FF0F172A')
    for columna in range(1, len(encabezados) + 1):
        celda = hoja.cell(row=fila_encabezado, column=columna)
        celda.font = Font(bold=True, color='FFFFFFFF')
        celda.fill = relleno
        celda.alignment = Alignment(horizontal='center')

    for fila in filas:
        hoja.append(list(fila))

    if totales:
        hoja.append([])
        for etiqueta, valor in totales:
            hoja.append([etiqueta, valor])
            hoja.cell(row=hoja.max_row, column=1).font = Font(bold=True)

    for columna in range(1, len(encabezados) + 1):
        ancho = max(
            [len(str(encabezados[columna - 1]))]
            + [len(str(fila[columna - 1])) for fila in filas if len(fila) >= columna]
            + [12]
        )
        hoja.column_dimensions[get_column_letter(columna)].width = min(ancho + 4, 45)

    hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=1)
    hoja.auto_filter.ref = f'A{fila_encabezado}:{get_column_letter(len(encabezados))}{hoja.max_row}'

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def construir_factura_pdf(factura: dict[str, Any], detalle: list[dict[str, Any]], venta: dict[str, Any]) -> bytes:
    """PDF de una factura de venta con datos del cliente, líneas y totales."""
    encabezados = ['Tipo', 'Descripción', 'Cantidad', 'Precio unitario', 'Descuento', 'Impuesto', 'Total']
    filas = [
        [
            (linea['tipo'] or '').capitalize(),
            linea['nombre'],
            f"{linea['cantidad']:g}",
            formato_moneda(linea['precioUnitario']),
            formato_moneda(linea['descuento']),
            formato_moneda(linea['impuesto']),
            formato_moneda(linea['total']),
        ]
        for linea in detalle
    ]
    subtitulo = (
        f"Factura {factura['numero']} · Fecha: {factura['fecha']}<br/>"
        f"Cliente: {factura['cliente']} · Documento: {factura['clienteDocumento'] or 'No registrado'}<br/>"
        f"Venta asociada: {venta.get('numero', '')} · Estado: {factura['estado']}"
    )
    totales = [
        ('Subtotal', formato_moneda(factura['subtotal'])),
        ('Descuentos', formato_moneda(factura['descuento'])),
        ('Impuestos', formato_moneda(factura['impuestos'])),
        ('Total a pagar', formato_moneda(factura['total'])),
    ]
    return construir_pdf(
        'Factura de venta',
        subtitulo,
        encabezados,
        filas,
        totales,
        notas=['Documento generado automáticamente por la plataforma SimonC.'],
    )
